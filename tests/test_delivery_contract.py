"""End-to-end synthetic workbook artifacts: explicit allowed edit vs realistic bad outputs."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sys import path as sys_path

sys_path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from verify_delivery import evaluate


class DeliveryContractFixtureTests(unittest.TestCase):
    def setUp(self):
        self.t = tempfile.TemporaryDirectory()
        self.root = Path(self.t.name)
        self.src = self.root / "original.xlsx"
        self.out = self.root / "result.xlsx"
        self.wb = Workbook()
        ws = self.wb.active
        ws.title = "Base"
        ws.append(["ID", "Amount", "CLASSIFICADO"])
        ws.append(["EXP-001", 10.0, "SIM"])
        ws.append(["EXP-002", 11.0, "NÃO"])
        ws["A1"].fill = PatternFill("solid", fgColor="B4C6E7")
        self.wb.save(self.src)

    def tearDown(self):
        self.t.cleanup()

    def contract(self, **opts):
        return dict(mode="EDIT_EXISTING", original=str(self.src), output=str(self.out),
                    header_rows={"Base": 1}, authorized_changed_ranges={"Base":["C2:C3"]},
                    allowed_status_values={"Base":{"CLASSIFICADO":["SIM","NÃO"]}}, **opts)

    def test_approved_existing_field_only(self):
        self.wb.active["C2"] = "NÃO"
        self.wb.save(self.out)
        result = evaluate(self.contract())
        self.assertTrue(result["pass"], result["errors"])

    def test_unrequested_extra_concilia_column_is_blocked(self):
        self.wb.active["D1"] = "CONCILIA"
        self.wb.save(self.out)
        self.assertFalse(evaluate(self.contract())["pass"])

    def test_unrequested_qa_sheet_is_blocked(self):
        self.wb.create_sheet("QA")
        self.wb.save(self.out)
        self.assertFalse(evaluate(self.contract())["pass"])

    def test_unrequested_green_status_fill_is_blocked(self):
        self.wb.active["C2"].fill = PatternFill("solid", fgColor="00FF00")
        self.wb.save(self.out)
        self.assertTrue(any("formatting" in e for e in evaluate(self.contract())["errors"]))

    def test_unrequested_revisar_status_is_blocked(self):
        self.wb.active["C2"] = "REVISAR"
        self.wb.save(self.out)
        self.assertFalse(evaluate(self.contract())["pass"])

    def test_unrequested_tolerance_sidebar_is_blocked(self):
        self.wb.active["F1"] = "Tolerância"
        self.wb.active["G1"] = .05
        self.wb.save(self.out)
        self.assertFalse(evaluate(self.contract())["pass"])

    def test_same_name_wrong_business_meaning_cannot_be_auto_verified(self):
        self.wb.save(self.out)
        result = evaluate(self.contract())
        self.assertIn("semantic_field_meaning", result["not_verified"])

    def test_missing_input_cannot_claim_pass(self):
        self.wb.save(self.out)
        data = self.contract()
        data["original"] = str(self.root / "missing.xlsx")
        with self.assertRaises(ValueError):
            evaluate(data)

    def test_explicit_added_header_at_fixed_position(self):
        self.wb.active["D1"] = "CONCILIA"
        self.wb.save(self.out)
        data = self.contract(authorized_new_headers={"Base":[{"column":4,"name":"CONCILIA"}]})
        result = evaluate(data)
        self.assertTrue(result["pass"], result["errors"])
        self.assertTrue(result["warnings"])


if __name__ == "__main__":
    unittest.main()
