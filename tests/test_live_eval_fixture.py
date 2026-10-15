"""Deterministic regression for opt-in real-agent artifact scorer."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from live_eval_fixture import prepare, score  # noqa: E402


class AgentArtifactScoringTests(unittest.TestCase):
    def test_correct_candidate_passes_both_gates(self):
        with tempfile.TemporaryDirectory() as root:
            folder = Path(root) / "task"
            prepare(folder)
            book = load_workbook(folder / "original.xlsx")
            book.active["C2"] = "NÃO"
            book.active["C3"] = "SIM"
            book.save(folder / "output.xlsx")
            result = score(folder)
            self.assertTrue(result["pass"], result)

    def test_wrong_values_fail_even_with_preserved_layout(self):
        with tempfile.TemporaryDirectory() as root:
            folder = Path(root) / "task"
            prepare(folder)
            book = load_workbook(folder / "original.xlsx")
            book.save(folder / "output.xlsx")
            result = score(folder)
            self.assertTrue(result["delivery_gate_pass"], result)
            self.assertFalse(result["correct_result_values"])
            self.assertFalse(result["pass"])

    def test_unapproved_column_fails_even_with_correct_values(self):
        with tempfile.TemporaryDirectory() as root:
            folder = Path(root) / "task"
            prepare(folder)
            book = load_workbook(folder / "original.xlsx")
            ws = book.active
            ws["C2"] = "NÃO"
            ws["C3"] = "SIM"
            ws["D1"] = "CONCILIA"
            book.save(folder / "output.xlsx")
            result = score(folder)
            self.assertTrue(result["correct_result_values"])
            self.assertFalse(result["delivery_gate_pass"])
            self.assertFalse(result["pass"])


if __name__ == "__main__":
    unittest.main()
