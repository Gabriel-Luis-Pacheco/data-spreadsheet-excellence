"""Prevents silently inventing delivery columns from an existing workbook."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from check_workbook_schema import compare_schema, schema_snapshot  # noqa: E402


class SchemaContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.src = Path(self.tmp.name) / "original.xlsx"
        self.out = Path(self.tmp.name) / "edited.xlsx"
        self.book = Workbook()
        ws = self.book.active
        ws.title = "B010"
        for i, title in enumerate(("Data", "Histórico", "Saldo", "CLASSIFICADO"), 1):
            ws.cell(3, i, title)
        ws["D4"] = "SIM"
        self.book.save(self.src)

    def tearDown(self):
        self.tmp.cleanup()

    def diff(self, additions=None):
        return compare_schema(
            schema_snapshot(self.src, {"B010": 3}),
            schema_snapshot(self.out, {"B010": 3}),
            allowed_headers=additions
        )

    def test_rejects_unrequested_concilia_qtd_embarque(self):
        ws = self.book["B010"]
        for i, title in enumerate(("CONCILIA", "QTD", "EMBARQUE"), 5):
            ws.cell(3, i, title)
        self.book.save(self.out)
        d = self.diff()
        self.assertFalse(d["ok"])
        self.assertEqual([v["after"] for v in d["header_differences"]],
                         ["CONCILIA", "QTD", "EMBARQUE"])
        self.assertTrue(d["column_growth_review"])

    def test_renaming_classificado_to_concilia_is_rejected(self):
        self.book["B010"]["D3"] = "CONCILIA"
        self.book.save(self.out)
        d = self.diff()
        self.assertFalse(d["ok"])
        self.assertEqual(d["header_differences"][0]["before"], "CLASSIFICADO")
        self.assertEqual(d["header_differences"][0]["after"], "CONCILIA")

    def test_populating_original_classificado_without_schema_change_passes(self):
        self.book["B010"]["D4"] = "NÃO"
        self.book.save(self.out)
        self.assertTrue(self.diff()["ok"])

    def test_explicit_new_column_is_whitelisted(self):
        self.book["B010"]["E3"] = "CONCILIA"
        self.book.save(self.out)
        self.assertTrue(self.diff({"B010": {"CONCILIA"}})["ok"])

    def test_unrequested_sheet_is_rejected(self):
        self.book.create_sheet("QA_Checks")
        self.book.save(self.out)
        d = self.diff()
        self.assertIn("QA_Checks", d["unapproved_added_sheets"])
        self.assertFalse(d["ok"])


if __name__ == "__main__":
    unittest.main()
