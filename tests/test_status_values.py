"""Focused tests for closed status vocabulary. No real Excel/COM needed."""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from check_status_values import check_status  # noqa: E402


class StatusContractTests(unittest.TestCase):
    def make_book(self, statuses):
        directory = tempfile.TemporaryDirectory()
        file = Path(directory.name) / "reconciliation.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "B010"
        for col, value in enumerate(["EMBARQUE", "Quantidade", "Saldo", "CONCILIA"], 1):
            sheet.cell(3, col, value)
        for row_num, (reference, status) in enumerate(statuses, 4):
            sheet.cell(row_num, 1, reference)
            if status is not None:
                sheet.cell(row_num, 4, status)
        book.save(file)
        return directory, file

    def test_unrequested_revisar_is_caught(self):
        temp, path = self.make_book([
            ("P0001", "SIM"), ("P0002", "REVISAR"), ("P0003", "NÃO")
        ])
        try:
            result = check_status(path, "B010", "CONCILIA", 3, ["SIM", "NÃO"], False)
            self.assertFalse(result["ok"])
            self.assertEqual(result["unexpected_count"], 1)
            self.assertEqual(result["status_counts"], {"NÃO": 1, "SIM": 1})
        finally:
            temp.cleanup()

    def test_blank_not_a_third_status_and_reporting_is_explicit(self):
        temp, path = self.make_book([
            ("P0001", "SIM"), ("AWB 123", None), ("P0003", "NÃO")
        ])
        try:
            result = check_status(path, "B010", "CONCILIA", 3, ["SIM", "NÃO"], False)
            self.assertTrue(result["ok"])
            self.assertEqual(result["blank_count"], 1)
            required = check_status(path, "B010", "CONCILIA", 3, ["SIM", "NÃO"], True)
            self.assertFalse(required["ok"])
        finally:
            temp.cleanup()

    def test_formulas_without_cached_results_are_not_falsely_certified(self):
        temp, path = self.make_book([("P0001", '=IF(TRUE,"SIM","NÃO")')])
        try:
            result = check_status(path, "B010", "CONCILIA", 3, ["SIM", "NÃO"], False)
            self.assertFalse(result["ok"])
            self.assertEqual(result["uncached_formula_count"], 1)
        finally:
            temp.cleanup()


if __name__ == "__main__":
    unittest.main()
