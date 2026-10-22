"""Tests for conservative human-copy auditing. Real model behavior needs live evals."""

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("audit_reader_copy", ROOT / "scripts" / "audit_reader_copy.py")
assert spec and spec.loader
audit = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = audit
spec.loader.exec_module(audit)


class ReaderCopyAuditTests(unittest.TestCase):
    def test_finds_high_confidence_filler(self):
        result = audit.audit_text("Neste relatório abordaremos uma análise robusta e abrangente.\n")
        rules = {row["rule"] for row in result}
        self.assertIn("empty_opening", rules)
        self.assertIn("promotional_filler", rules)

    def test_preserves_real_finance_domain_terms(self):
        text = "A EFD de setembro trouxe divergências no CST 060 e na NF-e 3526."
        self.assertEqual(audit.audit_text(text), [])

    def test_no_forced_global_dash_or_three_item_prohibition(self):
        self.assertEqual(audit.audit_text("Diferença de ICMS — três ajustes no período."), [])

    def test_skip_markdown_quote_and_code_examples(self):
        text = "> Neste relatório abordaremos nossos insights valiosos\n\n\x60\x60\x60\nEm conclusão, os resultados demonstram\n\x60\x60\x60\nConferência fiscal concluída."
        self.assertEqual(audit.audit_text(text, skip_examples=True), [])

    def test_excel_requires_explicit_author_range(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "working.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Relatório"
            ws["A1"] = "Ajustes de ICMS"
            ws["A2"] = "Em conclusão, os resultados demonstram..."
            ws["B2"] = "**Resumo**"
            ws["A3"] = "Descrições originais intactas"
            wb.save(path)
            with self.assertRaises(ValueError):
                audit.audit_file(path, None)
            issues = audit.audit_file(path, "Relatório!A1:B2")
            self.assertEqual({issue["rule"] for issue in issues}, {"empty_conclusion", "markdown_in_excel"})
            self.assertTrue(all("cell" in issue for issue in issues))
            self.assertEqual(audit.audit_file(path, "Relatório!A1:A1"), [])

    def test_no_formula_flags(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "test.xlsx"
            wb = openpyxl.Workbook()
            wb.active["A1"] = '="insights valiosos"'
            wb.save(path)
            self.assertEqual(audit.audit_file(path, "Sheet!A1:A1"), [])


if __name__ == "__main__":
    unittest.main()
