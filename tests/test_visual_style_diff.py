"""Targeted visual regression checks for the workbook comparator.

These are semantic openpyxl checks, not pixel-perfect Excel tests.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import workbook_diff  # noqa: E402


class VisualStyleDiffTests(unittest.TestCase):
    def test_added_green_sim_fill_is_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            before, after = Path(tmp) / "before.xlsx", Path(tmp) / "after.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "CONCILIA"
            ws["A2"] = "SIM"
            wb.save(before)
            ws["A2"].fill = PatternFill("solid", fgColor="C6EFCE")
            wb.save(after)

            diff = workbook_diff.compare(
                workbook_diff.workbook_snapshot(before, 1000),
                workbook_diff.workbook_snapshot(after, 1000),
            )
            self.assertIn("visual_cell_styles_changed", diff["risk_flags"])

    def test_change_to_existing_conditional_format_definition_is_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            before, after = Path(tmp) / "before.xlsx", Path(tmp) / "after.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["D4"] = "SIM"
            ws.conditional_formatting.add(
                "D4:D10",
                CellIsRule(operator="equal", formula=['"SIM"'],
                           fill=PatternFill("solid", fgColor="C6EFCE")),
            )
            wb.save(before)
            ws.conditional_formatting["D4:D10"][0].dxf.fill = PatternFill(
                "solid", fgColor="FFCC99"
            )
            wb.save(after)
            diff = workbook_diff.compare(
                workbook_diff.workbook_snapshot(before, 1000),
                workbook_diff.workbook_snapshot(after, 1000),
            )
            self.assertIn("conditional_formatting_details_changed", diff["risk_flags"])
            self.assertNotIn("conditional_formatting_changed", diff["risk_flags"])

    def test_same_styles_no_false_positive(self):
        with tempfile.TemporaryDirectory() as tmp:
            before, after = Path(tmp) / "before.xlsx", Path(tmp) / "after.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "CONCILIA"
            ws["A2"] = "SIM"
            ws["A2"].fill = PatternFill("solid", fgColor="C6EFCE")
            wb.save(before)
            wb.save(after)
            diff = workbook_diff.compare(
                workbook_diff.workbook_snapshot(before, 1000),
                workbook_diff.workbook_snapshot(after, 1000),
            )
            self.assertNotIn("visual_cell_styles_changed", diff["risk_flags"])
            self.assertNotIn("conditional_formatting_details_changed", diff["risk_flags"])


if __name__ == "__main__":
    unittest.main()
