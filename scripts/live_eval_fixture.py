#!/usr/bin/env python3
"""Prepare and score an opt-in REAL Codex run on an isolated Excel fixture.

Preparing/scoring is deterministic and CI-safe. This script NEVER launches an
agent or contacts a model: the user runs Codex separately in the fixture folder.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from verify_delivery import evaluate


def prepare(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / "original.xlsx"
    output = folder / "output.xlsx"
    if source.exists() or output.exists():
        raise ValueError("Fixture folder already contains original/output; use an EMPTY directory")
    book = Workbook()
    sheet = book.active
    sheet.title = "Base"
    sheet.append(["Lancamento", "Valor", "CLASSIFICADO"])
    sheet.append(["A1", 100.00, "SIM"])
    sheet.append(["A2", 120.00, "NÃO"])
    sheet["A1"].fill = __import__("openpyxl").styles.PatternFill(
        fill_type="solid", fgColor="B4C6E7")
    book.save(source)
    contract = {
        "mode": "EDIT_EXISTING",
        "original": str(source.resolve()),
        "output": str(output.resolve()),
        "header_rows": {"Base": 1},
        "authorized_changed_ranges": {"Base": ["C2:C3"]},
        "allowed_status_values": {"Base": {"CLASSIFICADO": ["SIM", "NÃO"]}},
    }
    (folder / "contract.json").write_text(json.dumps(contract, indent=2), encoding="utf-8")
    (folder / "prompt.txt").write_text(
        "Use the $data-spreadsheet-excellence skill. Edit ORIGINAL file original.xlsx "
        "into output.xlsx. In the existing CLASSIFICADO field, change only C2 to NÃO "
        "and C3 to SIM. Preserve everything else, including all original sheet names, "
        "headers, colors, filters and cell formats. Do not create any new column, tab, "
        "criteria block or execution report. Make the finished xlsx in this folder. "
        "Do not edit original.xlsx. Verify the result according to the skill.",
        encoding="utf-8",
    )
    return {"fixture_dir": str(folder), "original": str(source),
            "contract": str(folder / "contract.json"),
            "prompt": str(folder / "prompt.txt"), "expected_output": str(output)}


def score(folder: Path):
    contract = json.loads((folder / "contract.json").read_text(encoding="utf-8"))
    result = evaluate(contract)
    expected = {"C2": "NÃO", "C3": "SIM"}
    actual = {}
    if Path(contract["output"]).exists():
        book = load_workbook(contract["output"], read_only=True, data_only=True)
        try:
            actual = {coord: book["Base"][coord].value for coord in expected}
        finally:
            book.close()
    value_pass = actual == expected
    return {
        "pass": result["pass"] and value_pass,
        "delivery_gate_pass": result["pass"],
        "correct_result_values": value_pass,
        "expected": expected,
        "observed": actual,
        "gate_errors": result["errors"],
        "not_verified": result["not_verified"],
        "limitations": (
            "Synthetic file-based task only; cannot measure real source expertise, "
            "native Excel UI behavior, or future model reliability from one run."
        ),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="action", required=True)
    for name in ("prepare", "score"):
        sub.add_parser(name).add_argument("directory", type=Path)
    args = parser.parse_args()
    try:
        result = prepare(args.directory) if args.action == "prepare" else score(args.directory)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({"pass": False, "error": str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0 if args.action == "prepare" or result["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
