#!/usr/bin/env python3
"""Compare the ORIGINAL workbook's worksheet/header schema against an edited copy.

No guessing header rows: pass e.g. --header "Base:1" --header "B010:3".
Read-only; checks Excel headers/sheet inventory, not business semantics.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def header_at(ws, row: int) -> list[dict]:
    return [
        {"column": col, "name": str(cell.value)}
        for col in range(1, ws.max_column + 1)
        if (cell := ws.cell(row, col)).value is not None
        and str(cell.value).strip()
    ]


def schema_snapshot(path: Path, header_rows: dict[str, int]) -> dict:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            "sheet_names": list(book.sheetnames),
            "headers": {
                name: header_at(book[name], row)
                for name, row in header_rows.items()
                if name in book.sheetnames
            },
            "max_columns": {
                name: book[name].max_column for name in header_rows
                if name in book.sheetnames
            },
        }
    finally:
        book.close()


def compare_schema(original: dict, output: dict, allowed_sheets: set[str] | None = None,
                   allowed_headers: dict[str, set[str]] | None = None) -> dict:
    allowed_sheets = allowed_sheets or set()
    allowed_headers = allowed_headers or {}
    old_sheets = original["sheet_names"]
    new_sheets = output["sheet_names"]
    added_sheets = sorted(set(new_sheets) - set(old_sheets) - allowed_sheets)
    removed_sheets = sorted(set(old_sheets) - set(new_sheets))
    old_remaining = [x for x in old_sheets if x in new_sheets]
    new_remaining = [x for x in new_sheets if x in old_sheets]
    reordered = old_remaining != new_remaining
    differences = []
    for sheet, old_entries in original["headers"].items():
        if sheet not in output["headers"]:
            continue
        new_entries = output["headers"][sheet]
        prev = {x["column"]: x["name"] for x in old_entries}
        after = {x["column"]: x["name"] for x in new_entries}
        for col in sorted(set(prev) | set(after)):
            before = prev.get(col)
            now = after.get(col)
            if before == now:
                continue
            if before is None and now in allowed_headers.get(sheet, set()):
                continue
            differences.append({
                "sheet": sheet,
                "column": col,
                "before": before,
                "after": now,
                "type": "added" if before is None else ("removed" if now is None else "changed_or_moved"),
            })
    column_growth = [
        {"sheet": sheet, "original_max_column": max_column,
         "output_max_column": output["max_columns"][sheet]}
        for sheet, max_column in original["max_columns"].items()
        if sheet in output["max_columns"]
        and output["max_columns"][sheet] > max_column
    ]
    errors = []
    if added_sheets:
        errors.append("Unapproved added worksheet(s).")
    if removed_sheets:
        errors.append("Original worksheet(s) removed.")
    if reordered:
        errors.append("Original worksheet order changed.")
    if differences:
        errors.append("Unapproved header/schema changes.")
    return {
        "ok": not errors,
        "unapproved_added_sheets": added_sheets,
        "removed_sheets": removed_sheets,
        "original_sheet_order_changed": reordered,
        "header_differences": differences,
        "column_growth_review": column_growth,
        "errors": errors,
        "limitations": (
            "Structural header/sheet comparison only. Requires correct user-supplied "
            "header rows. Column growth may reflect styling and is review-only. "
            "Does not establish business meaning, record preservation, native "
            "Excel object functionality, or visual/pixel identity."
        ),
    }


def parse_sheet_row(value: str) -> tuple[str, int]:
    if ":" not in value:
        raise argparse.ArgumentTypeError(
            "Expected SHEET:ROW, e.g. Base:1 or B010:3."
        )
    sheet, number = value.rsplit(":", 1)
    if not sheet:
        raise argparse.ArgumentTypeError("Sheet name cannot be empty.")
    try:
        row = int(number)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Header row must be an integer.") from exc
    if row <= 0:
        raise argparse.ArgumentTypeError("Header row must be positive.")
    return sheet, row


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("original", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--header", action="append", type=parse_sheet_row, required=True,
                        help="Repeat SHEET:ROW for every in-scope sheet.")
    parser.add_argument("--allow-added-sheet", action="append", default=[])
    parser.add_argument("--allow-added-header", action="append", default=[],
                        help="SHEET:EXACT_HEADER_NAME, only for explicitly authorized additions.")
    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()
    rows = dict(args.header)
    if len(rows) != len(args.header):
        parser.error("Duplicate --header sheet.")
    additions: dict[str, set[str]] = {}
    for value in args.allow_added_header:
        if ":" not in value:
            parser.error("Use SHEET:EXACT_HEADER_NAME for --allow-added-header.")
        name, header = value.split(":", 1)
        if not name or not header:
            parser.error("Empty sheet/header in --allow-added-header.")
        additions.setdefault(name, set()).add(header)
    try:
        original = schema_snapshot(args.original, rows)
        edited = schema_snapshot(args.output, rows)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False),
              file=sys.stderr)
        return 2
    missing_header_sheets = [name for name in rows if name not in original["sheet_names"]]
    if missing_header_sheets:
        print(json.dumps({"ok": False, "error": "Sheet(s) absent from original",
                          "sheets": missing_header_sheets}, ensure_ascii=False),
              file=sys.stderr)
        return 2
    outcome = compare_schema(original, edited, set(args.allow_added_sheet), additions)
    print(json.dumps(outcome, ensure_ascii=False, indent=2 if args.pretty else None))
    return 0 if outcome["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
