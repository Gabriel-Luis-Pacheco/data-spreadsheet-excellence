#!/usr/bin/env python3
"""Read-only validation of allowed workbook status labels; not semantic reconciliation QA.

Example:
 python scripts/check_status_values.py book.xlsx --sheet B010 --header-row 3 \
   --column CONCILIA --allowed SIM --allowed NÃO
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def norm(value: object) -> str:
    return " ".join(str(value or "").split()).upper()


def check_status(
    file: Path,
    sheet: str,
    column_name: str,
    header_row: int,
    allowed: list[str],
    require_populated: bool,
) -> dict:
    wb = load_workbook(file, data_only=False, read_only=True)
    values_wb = load_workbook(file, data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet}")
        ws = wb[sheet]
        cached_ws = values_wb[sheet]
        target = norm(column_name)
        header_columns = [
            cell.column for cell in ws[header_row]
            if norm(cell.value) == target
        ]
        if len(header_columns) != 1:
            raise ValueError(
                f"Expected exactly one '{column_name}' column in row "
                f"{header_row}; found {len(header_columns)}."
            )

        col = header_columns[0]
        accepted = {norm(x) for x in allowed}
        if not accepted or "" in accepted:
            raise ValueError("At least one non-blank allowed status is required.")

        counts = Counter()
        unexpected: list[dict] = []
        blank_rows: list[int] = []
        unevaluated_formula_rows: list[int] = []
        for row in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row, col)
            value = cell.value
            if cell.data_type == "f":
                value = cached_ws.cell(row, col).value
                if value is None:
                    unevaluated_formula_rows.append(row)
                    continue
            normalized = norm(value)
            if not normalized:
                blank_rows.append(row)
                continue
            if normalized not in accepted:
                unexpected.append({"cell": cell.coordinate, "status": str(value)[:100]})
            else:
                counts[normalized] += 1

        errors: list[str] = []
        if unexpected:
            errors.append("Unexpected status values (closed vocabulary violated).")
        if unevaluated_formula_rows:
            errors.append("Formula result caches missing: Excel recalculation required.")
        if require_populated and blank_rows:
            errors.append("Blank status values found: check eligibility and intended population.")

        return {
            "sheet": sheet,
            "column": column_name,
            "allowed": sorted(accepted),
            "status_counts": dict(sorted(counts.items())),
            "unexpected_count": len(unexpected),
            "unexpected_examples": unexpected[:20],
            "blank_count": len(blank_rows),
            "blank_rows_sample": blank_rows[:20],
            "uncached_formula_count": len(unevaluated_formula_rows),
            "uncached_formula_rows_sample": unevaluated_formula_rows[:20],
            "ok": not errors,
            "errors": errors,
            "limitation": (
                "Checks vocabulary and available formula caches only; it cannot prove "
                "whether SIM/NÃO was assigned to the correct eligible shipment, or "
                "whether ignored records were excluded appropriately. Count/value "
                "bridge and key-field investigation are separately required."
            ),
        }
    finally:
        wb.close()
        values_wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--column", default="CONCILIA")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--allowed", action="append", default=None)
    parser.add_argument("--require-populated", action="store_true")
    args = parser.parse_args()
    if args.header_row < 1:
        parser.error("--header-row must be at least 1")
    try:
        data = check_status(
            args.workbook, args.sheet, args.column, args.header_row,
            args.allowed or ["SIM", "NÃO"], args.require_populated
        )
    except (OSError, ValueError, KeyError) as err:
        print(json.dumps({"ok": False, "error": str(err)}, ensure_ascii=False), file=sys.stderr)
        return 2
    print(json.dumps(data, indent=2, ensure_ascii=False))
    return 0 if data["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
