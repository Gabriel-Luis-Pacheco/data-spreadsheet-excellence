#!/usr/bin/env python3
"""Read-only semantic comparison for .xlsx/.xlsm workbooks.

This is a first-pass preservation/QA tool, not proof of full Excel fidelity.
It compares workbook/sheet structure and selected cell semantics without
executing macros, refreshing connections, or recalculating formulas.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

SUPPORTED = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def digest(items: list[str]) -> str:
    h = hashlib.sha256()
    for item in items:
        h.update(item.encode("utf-8", errors="replace"))
        h.update(b"\n")
    return h.hexdigest()


def safe_defined_names(wb) -> list[str]:
    names: list[str] = []
    try:
        for key in wb.defined_names:
            names.append(str(key))
    except Exception:
        pass
    return sorted(names)


def sheet_snapshot(ws, max_cells: int) -> dict[str, Any]:
    estimated = int(ws.max_row or 0) * int(ws.max_column or 0)
    scan_skipped = estimated > max_cells

    nonempty = 0
    formulas: list[str] = []
    content: list[str] = []
    number_formats: list[str] = []

    if not scan_skipped:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                nonempty += 1
                value_repr = repr(value)
                content.append(f"{cell.coordinate}\t{value_repr}")
                if isinstance(value, str) and value.startswith("="):
                    formulas.append(f"{cell.coordinate}\t{value}")
                number_formats.append(f"{cell.coordinate}\t{cell.number_format}")

    tables = []
    try:
        for name in sorted(ws.tables.keys()):
            table = ws.tables[name]
            tables.append({"name": name, "ref": str(table.ref)})
    except Exception:
        pass

    return {
        "state": ws.sheet_state,
        "max_row": int(ws.max_row or 0),
        "max_column": int(ws.max_column or 0),
        "estimated_rectangular_cells": estimated,
        "cell_scan_skipped": scan_skipped,
        "nonempty_cells": None if scan_skipped else nonempty,
        "formula_count": None if scan_skipped else len(formulas),
        "content_hash": None if scan_skipped else digest(content),
        "formula_hash": None if scan_skipped else digest(formulas),
        "number_format_hash": None if scan_skipped else digest(number_formats),
        "merged_ranges": sorted(str(x) for x in ws.merged_cells.ranges),
        "tables": tables,
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter and ws.auto_filter.ref else None,
        "data_validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
        "conditional_formatting_rules": len(ws.conditional_formatting),
        "charts": len(getattr(ws, "_charts", []) or []),
        "images": len(getattr(ws, "_images", []) or []),
        "protection_enabled": bool(ws.protection.sheet),
        "print_area": str(ws.print_area) if ws.print_area else None,
    }


def workbook_snapshot(path: Path, max_cells: int) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required: pip install openpyxl")

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    wb = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
        keep_links=True,
    )

    return {
        "path": str(path.resolve()),
        "extension": path.suffix.lower(),
        "sheet_order": [ws.title for ws in wb.worksheets],
        "defined_names": safe_defined_names(wb),
        "sheets": {ws.title: sheet_snapshot(ws, max_cells) for ws in wb.worksheets},
        "limitations": [
            "OOXML/openpyxl semantic comparison only.",
            "Does not execute or validate VBA/macros.",
            "Does not refresh Power Query/external connections.",
            "Does not recalculate formulas.",
            "Does not prove preservation of all drawings, pivots, embedded objects, or application-specific behavior.",
        ],
    }


def compare(before: dict[str, Any], after: dict[str, Any]) -> dict[str, Any]:
    changes: list[dict[str, Any]] = []
    risk_flags: list[str] = []

    def add(scope: str, field: str, a: Any, b: Any, risk: str | None = None) -> None:
        if a != b:
            changes.append({"scope": scope, "field": field, "before": a, "after": b})
            if risk:
                risk_flags.append(risk)

    add("workbook", "extension", before["extension"], after["extension"], "file_type_changed")
    add("workbook", "sheet_order", before["sheet_order"], after["sheet_order"], "sheet_structure_changed")
    add("workbook", "defined_names", before["defined_names"], after["defined_names"], "defined_names_changed")

    before_names = set(before["sheets"])
    after_names = set(after["sheets"])
    for name in sorted(before_names - after_names):
        changes.append({"scope": name, "field": "sheet_removed", "before": True, "after": False})
        risk_flags.append("sheet_removed")
    for name in sorted(after_names - before_names):
        changes.append({"scope": name, "field": "sheet_added", "before": False, "after": True})

    fields = [
        "state", "max_row", "max_column", "nonempty_cells", "formula_count",
        "content_hash", "formula_hash", "number_format_hash", "merged_ranges",
        "tables", "freeze_panes", "auto_filter", "data_validations",
        "conditional_formatting_rules", "charts", "images",
        "protection_enabled", "print_area",
    ]
    risky = {
        "state": "sheet_visibility_changed",
        "formula_count": "formula_structure_changed",
        "formula_hash": "formula_structure_changed",
        "merged_ranges": "merged_ranges_changed",
        "tables": "tables_changed",
        "data_validations": "data_validation_changed",
        "conditional_formatting_rules": "conditional_formatting_changed",
        "charts": "charts_changed",
        "images": "images_changed",
        "protection_enabled": "protection_changed",
    }

    for name in sorted(before_names & after_names):
        a = before["sheets"][name]
        b = after["sheets"][name]
        for field in fields:
            # If a full cell scan was skipped in either snapshot, hashes/counts
            # based on that scan are not comparable.
            if field in {"nonempty_cells", "formula_count", "content_hash", "formula_hash", "number_format_hash"}:
                if a["cell_scan_skipped"] or b["cell_scan_skipped"]:
                    continue
            add(name, field, a.get(field), b.get(field), risky.get(field))

    return {
        "before": before["path"],
        "after": after["path"],
        "change_count": len(changes),
        "risk_flags": sorted(set(risk_flags)),
        "changes": changes,
        "limitations": sorted(set(before["limitations"] + after["limitations"])),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("before")
    parser.add_argument("after")
    parser.add_argument("--max-cells", type=int, default=2_000_000)
    parser.add_argument("--pretty", action="store_true")
    parser.add_argument(
        "--fail-on-risk",
        action="store_true",
        help="Exit 2 when risk flags are detected. Useful as an opt-in regression gate.",
    )
    args = parser.parse_args()

    before_path = Path(args.before)
    after_path = Path(args.after)
    for path in (before_path, after_path):
        if not path.exists():
            raise SystemExit(f"File not found: {path}")
        if path.suffix.lower() not in SUPPORTED:
            raise SystemExit(f"Unsupported workbook format: {path.suffix}")

    result = compare(
        workbook_snapshot(before_path, args.max_cells),
        workbook_snapshot(after_path, args.max_cells),
    )
    print(json.dumps(result, indent=2 if args.pretty else None, ensure_ascii=False, default=str))
    if args.fail_on_risk and result["risk_flags"]:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
