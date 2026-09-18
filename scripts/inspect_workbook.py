#!/usr/bin/env python3
"""Read-only first-pass inventory for .xlsx/.xlsm workbooks.

Requires openpyxl. This script does not execute macros, refresh external data,
or prove full Excel fidelity. It highlights common workbook and agent-security
surfaces for follow-up review.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--pretty", action="store_true")
    parser.add_argument(
        "--max-cells",
        type=int,
        default=2_000_000,
        help="Skip full cell scan when a sheet's rectangular used range exceeds this budget.",
    )
    args = parser.parse_args()

    path = Path(args.workbook)
    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise SystemExit("Supported first-pass formats: .xlsx/.xlsm/.xltx/.xltm")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required: pip install openpyxl")

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    wb = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
        keep_links=True,
    )

    external_links = list(getattr(wb, "_external_links", []) or [])
    vba_archive_present = bool(getattr(wb, "vba_archive", None))

    result = {
        "path": str(path.resolve()),
        "extension": path.suffix.lower(),
        "keep_vba": keep_vba,
        "vba_archive_present": vba_archive_present,
        "external_links_count": len(external_links),
        "sheet_count": len(wb.worksheets),
        "sheets": [],
        "defined_names_count": len(wb.defined_names),
        "calculation": {},
        "risk_hints": [],
        "limitations": [
            "First-pass OOXML inventory only.",
            "Does not execute VBA/macros.",
            "Does not refresh Power Query/external connections.",
            "Does not fully inventory Power Query, embedded objects, threaded comments, or every drawing type.",
            "Does not prove preservation of every pivot/drawing/application-specific feature.",
            "Cell-level comments/hyperlinks are not scanned on sheets exceeding --max-cells.",
        ],
    }

    calc = getattr(wb, "calculation", None)
    if calc is not None:
        for attr in ("calcMode", "fullCalcOnLoad", "forceFullCalc", "calcId"):
            if hasattr(calc, attr):
                result["calculation"][attr] = getattr(calc, attr)

    hidden_sheets = 0
    total_comments = 0
    total_hyperlinks = 0

    for ws in wb.worksheets:
        estimated_cells = int(ws.max_row or 0) * int(ws.max_column or 0)
        scan_skipped = estimated_cells > args.max_cells

        formula_count = None
        hyperlink_count = None
        comment_count = None
        nonempty_count = None

        if not scan_skipped:
            formula_count = 0
            hyperlink_count = 0
            comment_count = 0
            nonempty_count = 0

            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        nonempty_count += 1
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formula_count += 1
                    if cell.hyperlink is not None:
                        hyperlink_count += 1
                    if cell.comment is not None:
                        comment_count += 1

            total_comments += comment_count
            total_hyperlinks += hyperlink_count

        if ws.sheet_state != "visible":
            hidden_sheets += 1

        result["sheets"].append({
            "title": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "estimated_rectangular_cells": estimated_cells,
            "cell_scan_skipped": scan_skipped,
            "nonempty_cells": nonempty_count,
            "formula_cells": formula_count,
            "comment_cells": comment_count,
            "merged_ranges": len(ws.merged_cells.ranges),
            "tables": sorted(ws.tables.keys()),
            "data_validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
            "conditional_formatting_rules": len(ws.conditional_formatting),
            "hyperlinks": hyperlink_count,
            "charts": len(getattr(ws, "_charts", []) or []),
            "images": len(getattr(ws, "_images", []) or []),
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter and ws.auto_filter.ref else None,
            "protection_enabled": bool(ws.protection.sheet),
            "print_area": str(ws.print_area) if ws.print_area else None,
        })

    if keep_vba or vba_archive_present:
        result["risk_hints"].append("macro_enabled_or_vba_container")
    if result["external_links_count"]:
        result["risk_hints"].append("external_links_present")
    if hidden_sheets:
        result["risk_hints"].append("hidden_or_very_hidden_sheets_present")
    if total_comments:
        result["risk_hints"].append("cell_comments_present")
    if total_hyperlinks:
        result["risk_hints"].append("hyperlinks_present")

    print(json.dumps(result, indent=2 if args.pretty else None, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
