#!/usr/bin/env python3
"""Fail-closed, read-only EDIT_EXISTING workbook delivery gate.

The allowlist is supplied from a trusted task instruction. This gate does NOT
verify the allowlist's authority, business semantics, or Excel F3 UI behavior.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from inspect_native_excel import inspect as inspect_native_parts


def _load(path):
    return load_workbook(path, data_only=False, read_only=False, keep_links=True,
                         keep_vba=path.suffix.lower() in {".xlsm", ".xltm"})


def _header(ws, row):
    return [(col, str(ws.cell(row, col).value))
            for col in range(1, ws.max_column + 1)
            if ws.cell(row, col).value is not None
            and str(ws.cell(row, col).value).strip()]


def _allowed_cell(row, col, ranges):
    for area in ranges:
        left, top, right, bottom = range_boundaries(area)
        if left <= col <= right and top <= row <= bottom:
            return True
    return False


def _style(cell):
    return (
        str(cell.font), str(cell.fill), str(cell.border), str(cell.alignment),
        cell.number_format, str(cell.protection),
    )


def _normalize(value):
    return " ".join(str(value).split()).upper() if value is not None else ""


def _sheet_structure(ws):
    """Stable, bounded comparison of user-visible/native container properties."""
    return {
        "visibility": ws.sheet_state,
        "freeze": str(ws.freeze_panes),
        "filter": str(ws.auto_filter.ref),
        "tables": [(name, str(ws.tables[name].ref)) for name in sorted(ws.tables)],
        "merged": sorted(str(r) for r in ws.merged_cells.ranges),
        "cf": sorted(
            f"{area.sqref}:{repr(rule)}:{repr(getattr(rule, 'dxf', None))}"
            for area in ws.conditional_formatting
            for rule in ws.conditional_formatting[area]
        ),
        "validations": repr(ws.data_validations),
        "print_area": str(ws.print_area),
        "chart_count": len(ws._charts),
        "image_count": len(ws._images),
    }


def evaluate(contract, max_cells=400_000):
    if contract.get("mode") != "EDIT_EXISTING":
        raise ValueError("This gate supports EDIT_EXISTING only. Never claim other modes were verified.")
    original = Path(contract["original"])
    output = Path(contract["output"])
    if original.resolve() == output.resolve():
        raise ValueError("Original and output must be separate files")
    if not original.is_file() or not output.is_file():
        raise ValueError("Original/output file not found")
    header_rows = contract["header_rows"]
    if not isinstance(header_rows, dict) or not header_rows:
        raise ValueError("Specify source header_rows by sheet")
    allowed_sheets = set(contract.get("authorized_new_sheets", []))
    allowed_new = contract.get("authorized_new_headers", {})
    allowed_cells = contract.get("authorized_changed_ranges", {})
    style_ranges = contract.get("authorized_style_ranges", {})
    object_changes = contract.get("authorized_object_changes", {})
    status_rules = contract.get("allowed_status_values", {})
    errors, warnings = [], []
    checked = 0
    before, after = _load(original), _load(output)
    try:
        a, b = set(after.sheetnames), set(before.sheetnames)
        errors.extend(f"Unapproved new sheet: {name}" for name in sorted(a-b-allowed_sheets))
        errors.extend(f"Original sheet removed: {name}" for name in sorted(b-a))
        if [x for x in after.sheetnames if x in b] != [x for x in before.sheetnames if x in a]:
            errors.append("Original sheet order changed")
        if set(before.sheetnames) != set(header_rows):
            raise ValueError("Header rows required for EVERY original worksheet, no partial verification")
        for sheet, row in header_rows.items():
            if not isinstance(row, int) or row < 1:
                raise ValueError(f"Bad header row for {sheet}")
        if any(set(cfg) - b for cfg in (allowed_new, allowed_cells, style_ranges, object_changes, status_rules)):
            raise ValueError("Contract scopes unknown original worksheets")
        for sheet in before.sheetnames:
            if sheet not in a:
                continue
            ws, dst = before[sheet], after[sheet]
            hdr = header_rows[sheet]
            old, new = _header(ws, hdr), _header(dst, hdr)
            additions = allowed_new.get(sheet, [])
            if not isinstance(additions, list):
                raise ValueError(f"authorized_new_headers.{sheet} must be a list")
            approved: dict[int, str] = {}
            for item in additions:
                col = item.get("column")
                if isinstance(col, str):
                    col = column_index_from_string(col)
                if not isinstance(col, int) or col < 1:
                    raise ValueError("Authorized header needs exact Excel column index")
                if col in approved:
                    raise ValueError(f"Duplicated authorized header position {sheet}!{col}")
                approved[col] = str(item.get("name", ""))
                if (col, approved[col]) not in new:
                    errors.append(f"Authorized header absent/wrong position: {sheet}!{col}")
            expected = [name for _, name in old]
            actual = [name for col, name in new if col not in approved]
            if expected != actual:
                errors.append(f"Unexpected header/order changes in {sheet}: {expected!r} => {actual!r}")
            if len(old) + len(additions) != len(new):
                errors.append(f"Header count changed beyond permission on {sheet}")

            ranges = allowed_cells.get(sheet, [])
            styles = style_ranges.get(sheet, [])
            for area in ranges + styles:
                try:
                    range_boundaries(area)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"Bad authorized range {sheet}!{area}") from exc
            size = max(ws.max_row, dst.max_row) * max(ws.max_column, dst.max_column)
            if size > max_cells:
                raise ValueError(f"{sheet}: {size} cells exceeds scan cap {max_cells}; NOT VERIFIED")
            allowed_property_names = set(object_changes.get(sheet, []))
            old_structure, new_structure = _sheet_structure(ws), _sheet_structure(dst)
            for key, old_value in old_structure.items():
                if old_value != new_structure[key] and key not in allowed_property_names:
                    errors.append(f"Unapproved Excel object/layout change: {sheet}.{key}")

            # Map destination columns back to original columns after *authorized*
            # insertions; do NOT skip cell validation because an insertion exists.
            added_columns = sorted(approved)
            for row in range(1, max(ws.max_row, dst.max_row) + 1):
                for dest_col in range(1, max(ws.max_column + len(approved), dst.max_column) + 1):
                    if dest_col in approved:
                        continue  # Entire new field is in the user's explicit allowed additions.
                    original_col = dest_col - sum(index < dest_col for index in added_columns)
                    if original_col < 1:
                        continue
                    source_cell = ws.cell(row, original_col)
                    result_cell = dst.cell(row, dest_col)
                    checked += 1
                    if source_cell.value != result_cell.value and not _allowed_cell(row, original_col, ranges):
                        errors.append(f"Unauthorized cell change: {sheet}!{result_cell.coordinate}")
                    if _style(source_cell) != _style(result_cell) and not _allowed_cell(row, original_col, styles):
                        errors.append(f"Unapproved formatting change: {sheet}!{result_cell.coordinate}")
                    if len(errors) >= 40:
                        break
                if len(errors) >= 40:
                    break
            for header, permitted in status_rules.get(sheet, {}).items():
                hits = [col for col, name in new if name == header]
                if len(hits) != 1:
                    errors.append(f"Status field missing/ambiguous: {sheet}.{header}")
                    continue
                allowed = {_normalize(x) for x in permitted}
                if not allowed or "" in allowed:
                    raise ValueError(f"Invalid status allowlist: {sheet}.{header}")
                formula_count = 0
                for idx in range(hdr + 1, dst.max_row + 1):
                    cell = dst.cell(idx, hits[0])
                    if cell.data_type == "f":
                        formula_count += 1
                        continue
                    if _normalize(cell.value) and _normalize(cell.value) not in allowed:
                        errors.append(f"Unexpected status {sheet}!{cell.coordinate}: {cell.value!r}")
                        if len(errors) >= 40:
                            break
                if formula_count:
                    warnings.append(f"{sheet}.{header}: {formula_count} formula results need independent Excel recalc verification")
        # OOXML check: native part preservation is necessary, never sufficient.
        parts_before, parts_after = inspect_native_parts(original), inspect_native_parts(output)
        key_map = {
            "pivot_table": "verified_pivot_structure_count",
            "excel_table": "linked_excel_table_count",
            "chart": "chart_parts",
            "slicer": "slicer_parts",
            "connection": "connections_part",
        }
        for title, attr in key_map.items():
            original_count = parts_before["objects"].get(attr, 0)
            actual_count = parts_after["objects"].get(attr, 0)
            if actual_count < original_count:
                errors.append(f"Native Excel object lost: {title} {original_count} => {actual_count}")
        for native in contract.get("required_native_objects", []):
            if native not in key_map:
                raise ValueError(f"Unsupported required_native_objects entry: {native}")
            if parts_after["objects"].get(key_map[native], 0) == 0:
                errors.append(f"Required native Excel structure absent: {native}")
        if contract.get("required_native_objects"):
            warnings.append("OOXML native object structure found, but Excel UI/refresh remains unverified")
        if allowed_sheets:
            warnings.append("Authorized new worksheets not compared for cell-level scope")
        return {
            "pass": not errors,
            "checked_cells": checked,
            "errors": errors[:40],
            "warnings": sorted(set(warnings)),
            "verified": ["source_sheet_inventory", "original_header_order_and_approved_position",
                         "unchanged_cells_and_styles_including_insertions", "sheet_objects",
                         "native_object_part_preservation",
                         "literal_closed_status_values"],
            "not_verified": ["semantic_field_meaning", "Excel_native_F3_interactivity",
                             "calculated_formula_results", "authorization_of_contract_permissions"],
        }
    finally:
        before.close()
        after.close()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--contract", required=True, type=Path)
    p.add_argument("--max-cells", type=int, default=400_000)
    p.add_argument("--pretty", action="store_true")
    a = p.parse_args()
    try:
        contract = json.loads(a.contract.read_text(encoding="utf-8"))
        result = evaluate(contract, max_cells=a.max_cells)
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(json.dumps({"pass": False, "error": str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2 if a.pretty else None))
    return 0 if result["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
